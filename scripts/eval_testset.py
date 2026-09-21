"""Pipeline-Ergebnis gegen eine Loesungsdatei messen (Klasse UND Merkmale).

    python scripts/eval_testset.py out/lts30 data/testsets/LTS_Testset_30_LOESUNG.xlsx
    python scripts/eval_testset.py out/lts30 <loesung.xlsx> --write-reference   # reference.json fuer compare

Loesungsdatei: Blatt 1 = Artikelnummer | … | ETIM-Klasse (Spalte "ETIM-Klasse"),
Blatt "Loesung_Merkmale" = Artikelnummer | ETIM-Klasse | Merkmal (EF) | Wert 1 | Wert 2 | Wert-Details | Einheit.
Werte mit Details "NA" gelten als "nicht anwendbar" und zaehlen nicht als Soll.

Gemessen wird, was ein Grosshaendler sieht:
- Klassen-Trefferquote gesamt und je Soll-Klasse; wie viele Fehler die Review-Queue auffaengt
- Merkmale nur bei richtiger Klasse (sonst vergleicht man Aepfel mit Birnen):
  Praezision (gefuellt UND richtig / gefuellt), Abdeckung (richtig / Soll),
  und FALSCH GEFUELLT — der teuerste Fehler, weil er ohne Pruefung beim Kunden landet.
Ergebnis: <job>/eval.md und <job>/eval.json.
"""
from __future__ import annotations

import argparse
import json
import sys
from collections import Counter, defaultdict
from pathlib import Path

import openpyxl

NUM_TOL = 0.01  # 1 % relative Toleranz bei Zahlen (Rundung im Katalog)


def _s(v) -> str:
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    return str(v).strip()


def load_solution(path: Path):
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    ws = wb.worksheets[0]
    it = ws.iter_rows(values_only=True)
    head = [_s(h) for h in next(it)]
    ci, ki = head.index("Artikelnummer"), head.index("ETIM-Klasse")
    classes = {_s(r[ci]): _s(r[ki]) for r in it if r and r[ci] is not None}
    feats: dict[str, dict[str, tuple[str, str]]] = defaultdict(dict)
    if "Loesung_Merkmale" in wb.sheetnames:
        it = wb["Loesung_Merkmale"].iter_rows(values_only=True)
        next(it)
        for r in it:
            pid, _, fid, v1, v2, det = (_s(x) for x in r[:6])
            if pid and fid and det != "NA" and v1 not in ("", "-"):
                feats[pid][fid] = (v1, v2)
    return classes, feats


def _num(x: str):
    try:
        return float(str(x).replace(",", "."))
    except (TypeError, ValueError):
        return None


def same(truth: str, pred: str | None) -> bool:
    if pred is None or pred == "":
        return False
    t, p = truth.strip().lower(), str(pred).strip().lower()
    if t == p:
        return True
    tn, pn = _num(t), _num(p)
    if tn is not None and pn is not None:
        return abs(tn - pn) <= max(abs(tn) * NUM_TOL, 1e-9)
    return False


def evaluate(job: Path, solution: Path) -> dict:
    classes, sol_feats = load_solution(solution)
    classified = json.loads((job / "classified.json").read_text())
    enriched_p = job / "enriched.json"
    enriched = {e["product"]["supplier_pid"]: e for e in json.loads(enriched_p.read_text())} if enriched_p.exists() else {}

    per_class = defaultdict(lambda: [0, 0])  # soll-klasse -> [richtig, gesamt]
    rows, confusions = [], Counter()
    hits = n = none = review = wrong_caught = wrong_total = 0
    simulated = False
    f_truth = f_filled = f_right = f_wrong = 0
    f_by_feature = defaultdict(lambda: [0, 0, 0])  # fid -> [soll, richtig, falsch]
    model_name = ""
    for c in classified:
        pid = c["product"]["supplier_pid"]
        if pid not in classes:
            continue
        simulated |= bool(c.get("simulated"))
        model_name = c.get("model", model_name)
        truth, pred = classes[pid], c["decision"].get("class_id")
        ok = pred == truth
        n += 1
        hits += ok
        none += pred is None
        review += bool(c.get("needs_review"))
        per_class[truth][0] += ok
        per_class[truth][1] += 1
        if not ok:
            wrong_total += 1
            wrong_caught += bool(c.get("needs_review"))
            confusions[(truth, pred or "—")] += 1
        rank = next((i + 1 for i, k in enumerate(c.get("candidates", [])) if k["class_id"] == truth), None)
        row = {"pid": pid, "name": c["product"]["name"][:70], "soll": truth, "ist": pred,
               "ok": ok, "review": c.get("needs_review"), "conf": c["decision"].get("confidence")}
        # Merkmale nur bei richtiger Klasse
        e = enriched.get(pid)
        if ok and e is not None:
            pv = {f["feature_id"]: f for f in e["features"]}
            soll = sol_feats.get(pid, {})
            r_ok = r_bad = 0
            for fid, (v1, v2) in soll.items():
                f_truth += 1
                f_by_feature[fid][0] += 1
                fv = pv.get(fid)
                if not fv or fv.get("value") in (None, ""):
                    continue
                good = same(v1, fv["value"]) and (not v2 or v2 == v1 or same(v2, fv.get("value_max") or fv["value"]))
                r_ok += good
                r_bad += not good
                f_by_feature[fid][1 if good else 2] += 1
            # gefuellt, obwohl die Loesung dort nichts (oder NA) hat -> auch falsch gefuellt
            extra = sum(1 for fid, fv in pv.items() if fv.get("value") not in (None, "") and fid not in soll)
            f_right += r_ok
            f_wrong += r_bad + extra
            f_filled += r_ok + r_bad + extra
            row.update(merkmale_soll=len(soll), merkmale_richtig=r_ok, merkmale_falsch=r_bad + extra)
        rows.append(row)
        row["rang_soll_in_kandidaten"] = rank

    res = {
        "job": job.name, "modell": model_name, "simuliert": simulated, "artikel": n,
        "klasse_treffer": hits, "klasse_quote": round(hits / n, 4) if n else None,
        "ohne_klasse": none, "zur_pruefung": review,
        "fehler_von_review_aufgefangen": f"{wrong_caught}/{wrong_total}",
        "merkmale": {"soll": f_truth, "gefuellt": f_filled, "richtig": f_right, "falsch_gefuellt": f_wrong,
                     "praezision": round(f_right / f_filled, 4) if f_filled else None,
                     "abdeckung": round(f_right / f_truth, 4) if f_truth else None},
        "je_klasse": {k: {"richtig": v[0], "gesamt": v[1]} for k, v in sorted(per_class.items())},
        "verwechslungen": [{"soll": a, "ist": b, "anzahl": k} for (a, b), k in confusions.most_common()],
        "schlechteste_merkmale": sorted(
            ({"merkmal": k, "soll": v[0], "richtig": v[1], "falsch": v[2]} for k, v in f_by_feature.items()),
            key=lambda x: (x["richtig"] / x["soll"]) if x["soll"] else 1)[:15],
        "artikel_detail": rows,
    }
    return res


def to_md(r: dict) -> str:
    m = r["merkmale"]
    pct = lambda x: "—" if x is None else f"{x:.0%}"
    out = [f"# Messung {r['job']} · Modell {r['modell']}" + (" · **SIMULIERT, keine echte Zahl**" if r["simuliert"] else ""),
           "", f"- Klassen richtig: **{r['klasse_treffer']}/{r['artikel']} ({pct(r['klasse_quote'])})**",
           f"- Ohne Klasse: {r['ohne_klasse']} · zur Prüfung: {r['zur_pruefung']} · "
           f"Fehler von der Prüfung aufgefangen: {r['fehler_von_review_aufgefangen']}",
           f"- Merkmale (nur richtig klassifizierte): Präzision **{pct(m['praezision'])}**, "
           f"Abdeckung **{pct(m['abdeckung'])}**, falsch gefüllt **{m['falsch_gefuellt']}** von {m['gefuellt']}",
           "", "## Je Klasse", "", "| Soll-Klasse | richtig |", "|---|---|"]
    out += [f"| {k} | {v['richtig']}/{v['gesamt']} |" for k, v in r["je_klasse"].items()]
    if r["verwechslungen"]:
        out += ["", "## Verwechslungen", "", "| Soll | Ist | Anzahl |", "|---|---|---|"]
        out += [f"| {v['soll']} | {v['ist']} | {v['anzahl']} |" for v in r["verwechslungen"]]
    if r["schlechteste_merkmale"]:
        out += ["", "## Schwächste Merkmale", "", "| Merkmal | Soll | richtig | falsch |", "|---|---|---|---|"]
        out += [f"| {v['merkmal']} | {v['soll']} | {v['richtig']} | {v['falsch']} |" for v in r["schlechteste_merkmale"]]
    return "\n".join(out) + "\n"


def main(argv=None) -> int:
    ap = argparse.ArgumentParser()
    ap.add_argument("job", type=Path)
    ap.add_argument("solution", type=Path)
    ap.add_argument("--write-reference", action="store_true", help="reference.json fuer compare/Cockpit schreiben")
    a = ap.parse_args(argv)
    if a.write_reference:
        classes, _ = load_solution(a.solution)
        a.job.mkdir(parents=True, exist_ok=True)
        (a.job / "reference.json").write_text(json.dumps({"reference": classes}, indent=2))
        print(f"reference.json: {len(classes)} Artikel → {a.job / 'reference.json'}")
        if not (a.job / "classified.json").exists():
            return 0
    r = evaluate(a.job, a.solution)
    (a.job / "eval.json").write_text(json.dumps(r, ensure_ascii=False, indent=2))
    (a.job / "eval.md").write_text(to_md(r))
    print(to_md(r).split("## Je Klasse")[0])
    return 0


if __name__ == "__main__":
    sys.exit(main())
